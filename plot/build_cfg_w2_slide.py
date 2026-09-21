#!/usr/bin/env python3
"""Build + append slide-354-w2: EE substrate validity, but the CFG run scored at
w=2 (others at w=0, guidance N/A for non-CFG). Config-derived attrs. PowerPoint CLOSED."""
import datetime, shutil
from copy import deepcopy
from pathlib import Path
import numpy as np, pandas as pd
import matplotlib; matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import Patch
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.enum.text import MSO_AUTO_SIZE

ROOT = "/Users/soldatmat/Documents/terpene_synthases"
D = Path(f"{ROOT}/dplm/run/class_predictor/slide306_eval/ee_substrate_metric_2026-06-15")
PPTX = f"{ROOT}/presentation/dplm.pptx"
FIG = D / "ee_substrate_by_arch_w2.png"
CFG_W2_SUB = 0.863  # A3cfg-PRE class-0 P(FPP) at w=2 (w0 was ~0.815/0.840)
ARCH_COLOR = {"prepend": "#4C72B0", "CA": "#DD8452", "miniCA": "#55A868"}
COND_EDGE = {"normal": "#333333", "cfg": "#E6B800", "neighbor": "#C0249B"}
COND_LW = {"normal": 0.5, "cfg": 2.4, "neighbor": 2.4}


def short(run):
    s = run.replace("TPS_dplm_150m_class_", "").replace("first_cyclization_", "")
    for tok in ("_lr1em3","_lr5em4","_lr1em4","_lr2p5em4","_lr1em2","_wu2000","_wu4000"):
        i = s.find(tok)
        if i > 0: s = s[:i]
    return s[:42]


def build_fig():
    df = pd.read_csv(D/"ee_substrate_scores.csv")
    attrs = pd.read_csv(D/"run_attributes_from_config.csv")[["run","arch","cond_type","embed_state"]]
    d = df[df["class"].notna() & ~df["run"].str.contains("smoke")].merge(attrs, on="run")
    # override the CFG run to its w=2 substrate validity
    cfgmask = d["cond_type"] == "cfg"
    d.loc[cfgmask, "correct_substrate_score"] = CFG_W2_SUB
    d.loc[cfgmask, "run"] = d.loc[cfgmask, "run"] + "  @w2"
    d = d.sort_values("correct_substrate_score").reset_index(drop=True)
    uncond = df[df["class"].isna()]["correct_substrate_score"].mean()
    cm = d["correct_substrate_score"].mean()
    n = len(d)
    fig, ax = plt.subplots(figsize=(12.5, max(8.0, 0.20*n)), constrained_layout=True)
    y = np.arange(n)
    for yi, (_, r) in zip(y, d.iterrows()):
        ax.barh(yi, r["correct_substrate_score"], height=0.78, color=ARCH_COLOR.get(r["arch"],"0.6"),
                hatch=("///" if r["embed_state"]=="unfrozen" else None),
                edgecolor=COND_EDGE[r["cond_type"]], linewidth=COND_LW[r["cond_type"]], zorder=3)
    ax.axvline(cm, color="0.35", ls="--", lw=1); ax.axvline(uncond, color="0.35", ls=":", lw=1)
    ax.set_yticks(y); ax.set_yticklabels([short(r) for r in d["run"]], fontsize=6.5)
    ax.set_ylim(-0.7, n+0.7); ax.set_xlim(0,1.0)
    ax.set_xlabel("EE P(FPP) [class-0 sesqui] — CFG run at w=2, others w=0", fontsize=10)
    arch_h=[Patch(facecolor=ARCH_COLOR[a],edgecolor="0.3",label=a) for a in ["prepend","CA","miniCA"]]
    embed_h=[Patch(facecolor="0.8",edgecolor="0.3",label="frozen"),Patch(facecolor="0.8",edgecolor="0.3",hatch="///",label="unfrozen")]
    cond_h=[Patch(facecolor="white",edgecolor=COND_EDGE[c],linewidth=COND_LW[c],label=c) for c in ["normal","cfg","neighbor"]]
    l1=ax.legend(handles=arch_h,title="architecture",loc="lower right",bbox_to_anchor=(1.0,0.02),fontsize=8,title_fontsize=8)
    l2=ax.legend(handles=embed_h,title="embedding",loc="lower right",bbox_to_anchor=(1.0,0.26),fontsize=8,title_fontsize=8)
    ax.legend(handles=cond_h,title="conditioning",loc="lower right",bbox_to_anchor=(1.0,0.48),fontsize=8,title_fontsize=8)
    ax.add_artist(l1); ax.add_artist(l2)
    fig.suptitle("EE substrate validity — CFG run at w=2 (others w=0; guidance N/A for non-CFG)",
                 fontsize=12, fontweight="bold")
    fig.savefig(FIG, dpi=200, bbox_inches="tight"); plt.close(fig)
    print(f"fig: {FIG}  (CFG bar @ {CFG_W2_SUB})")


def fit(img, slot):
    sl,st,sw,sh=slot
    with Image.open(img) as im: iw,ih=im.size
    r,sr=iw/ih,sw/sh; w,h=(sw,sw/r) if r>sr else (sh*r,sh)
    return sl+(sw-w)/2, st+(sh-h)/2, w, h

def settf(tf,t):
    p0=tf.paragraphs[0]
    if p0.runs:
        p0.runs[0].text=t
        for e in p0.runs[1:]: e._r.getparent().remove(e._r)
    else: p0.text=t
    for ep in list(tf.paragraphs[1:]): ep._p.getparent().remove(ep._p)

def setann(shape,text):
    tf=shape.text_frame; tf.word_wrap=True; tf.auto_size=MSO_AUTO_SIZE.NONE
    ls=text.split("\n"); tf.paragraphs[0].text=ls[0]
    for ep in list(tf.paragraphs[1:]): ep._p.getparent().remove(ep._p)
    for ln in ls[1:]: tf.add_paragraph().text=ln
    for p in tf.paragraphs:
        for r in p.runs: r.font.size=Pt(10)
    shape.height=Inches(min(len(ls)*10*1.2/72+0.15, 7.4-Emu(shape.top).inches))

def append():
    if list(Path(PPTX).parent.glob("~$*.pptx")): raise SystemExit("ABORT: PowerPoint open.")
    shutil.copy2(PPTX, f"/tmp/dplm_pptx_backup_{datetime.datetime.now():%Y%m%d_%H%M%S}.pptx")
    prs=Presentation(PPTX); tmpl=prs.slides[232]
    t=prs.slides.add_slide(prs.slide_layouts[1])
    for ph in list(t.placeholders): ph.element.getparent().remove(ph.element)
    for sh in tmpl.shapes:
        if sh.shape_type==13: continue
        t.shapes._spTree.insert_element_before(deepcopy(sh.element),"p:extLst")
    l,tp,w,h=fit(FIG,(0.12,0.96,8.48,6.35)); t.shapes.add_picture(str(FIG),Inches(l),Inches(tp),Inches(w),Inches(h))
    ann=("Same as slide #354 but the CFG run is scored at its\n"
         "deployable w=2 (others stay w=0 — guidance only\n"
         "applies to CFG-trained models).\n\n"
         "A3cfg-PRE class-0 substrate: w0 0.84 -> w2 0.86\n"
         "(+0.02); isTPS flat ~0.92. So at w=2 it stays the\n"
         "top bar, modestly above its w=0 value.\n\n"
         "Still class-0 sesqui, sequence-only EE. w2 scale\n"
         "verified vs slide #354 (w0 reproduced ~0.82).")
    for sh in t.shapes:
        if sh.shape_type==13 or not sh.has_text_frame: continue
        if sh.name=="TextovéPole 11": settf(sh.text_frame,"EE substrate validity — CFG run at w=2 (others w=0)")
        elif sh.name=="TextovéPole 12": settf(sh.text_frame,"")
        elif sh.name=="TextBox 5": setann(sh,ann)
    prs.save(PPTX)
    print(f"appended CFG-w2 slide at #{len(prs.slides)}; total {len(prs.slides)}")


if __name__ == "__main__":
    build_fig(); append()
