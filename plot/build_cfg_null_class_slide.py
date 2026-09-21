#!/usr/bin/env python3
"""Append a simple, minimalistic CFG slide to dplm.pptx that VISUALIZES the
learned-null-class idea: a two-step left->right schematic — (1) training learns a
null embedding via condition dropout, (2) inference steers with w along the
null->conditional->guided direction. Conceptual diagram, NOT a metrics plot.
PowerPoint must be CLOSED. Run with /opt/miniconda3/bin/python3.

Modeled on build_cfg_explainer_slide.py — reuses its exact append constants and
fit()/set_title()/set_annot() helpers and /tmp backup discipline."""
import datetime, shutil
from copy import deepcopy
from pathlib import Path
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.enum.text import MSO_AUTO_SIZE

ROOT = "/Users/soldatmat/Documents/terpene_synthases"
PPTX = f"{ROOT}/presentation/dplm.pptx"
FIG = f"{ROOT}/dplm/plot/cfg_null_class.png"
TEMPLATE_SLIDE_INDEX = 232; LAYOUT_INDEX = 1
TITLE_SHAPE = "TextovéPole 11"; NUM_SHAPE = "TextovéPole 12"; ANNOT_SHAPE = "TextBox 5"
PIC_SLOT = (0.12, 0.96, 8.48, 6.35); ANNOT_PT = 10.0

COND = "#4C72B0"    # conditional (target class)
NULL = "#888888"    # unconditional (learned null)
GOLD = "#E6B800"    # the w knob / guidance accent — matches the other CFG slides

TITLE = "CFG — a learned null class for controllable conditioning"
ANNOT = ("Conditioning: class label →\n"
         "learned embedding, prepended\n"
         "as a token to frozen ESM-2 DPLM.\n\n"
         "Condition dropout (training):\n"
         "15% of the time the class is\n"
         "swapped for ONE learned null ∅.\n"
         "→ one network learns both\n"
         "p(x | c) and p(x | ∅).\n\n"
         "Steer at inference:\n"
         "∅ is a learned baseline; the\n"
         "c − ∅ gap is a direction.\n"
         "w sets how far to push:\n"
         "• w = 0 → plain conditional\n"
         "• w ↑ → stronger class")


def _box(ax, cx, cy, w, h, label, edge, face, fontsize=13, fontweight="bold",
         textcolor=None):
    ax.add_patch(FancyBboxPatch((cx - w / 2, cy - h / 2), w, h,
                                boxstyle="round,pad=0.006,rounding_size=0.012",
                                linewidth=1.8, edgecolor=edge, facecolor=face))
    ax.text(cx, cy, label, ha="center", va="center", fontsize=fontsize,
            color=textcolor or edge, fontweight=fontweight)


def build_fig():
    fig = plt.figure(figsize=(9.5, 6.0))
    ax = fig.add_axes([0, 0, 1, 1]); ax.axis("off")
    ax.set_xlim(0, 1); ax.set_ylim(0, 1)

    # ---- panel divider ------------------------------------------------------
    ax.plot([0.5, 0.5], [0.06, 0.9], color="0.85", lw=1.0, zorder=0)

    # ================= LEFT PANEL: TRAINING — learn a null ===================
    ax.text(0.25, 0.95, "Training — learn a null", ha="center", va="top",
            fontsize=14, color="0.25", fontweight="bold")

    # row of class-token boxes
    bw, bh = 0.085, 0.085
    xs = [0.085, 0.185, 0.285]
    labels = [r"$c_0$", r"$c_1$", r"$c_2$"]
    ytok = 0.72
    for x, lab in zip(xs, labels):
        _box(ax, x, ytok, bw, bh, lab, COND, "#EAF0F8")
    ax.text(0.355, ytok, r"$\cdots$", ha="center", va="center",
            fontsize=15, color=COND)
    # the distinct null box, set slightly apart
    _box(ax, 0.43, ytok, bw, bh, r"$\varnothing$", NULL, "#EFEFEF", fontsize=16)
    ax.text(0.43, ytok - 0.075, "null", ha="center", va="top",
            fontsize=9.5, color=NULL)
    ax.text(0.185, ytok + 0.075, "class tokens", ha="center", va="bottom",
            fontsize=9.5, color=COND)

    # dropout arrow: a class can be replaced by the null
    ax.annotate("", xy=(0.43, ytok - 0.005), xytext=(0.285, ytok - 0.005),
                arrowprops=dict(arrowstyle="-|>", color=GOLD, lw=2.0,
                                connectionstyle="arc3,rad=-0.45"))
    ax.text(0.36, ytok - 0.155, "15% of the time → ∅", ha="center", va="center",
            fontsize=11, color=GOLD, fontweight="bold")
    ax.text(0.36, ytok - 0.205, "(condition dropout)", ha="center", va="center",
            fontsize=9.5, color="0.5", style="italic")

    # downward flow into the one network
    ax.annotate("", xy=(0.25, 0.405), xytext=(0.25, 0.5),
                arrowprops=dict(arrowstyle="-|>", color="0.5", lw=1.8))
    _box(ax, 0.25, 0.335, 0.30, 0.10, "one frozen ESM-2 DPLM", "0.35",
         "#F6F6F6", fontsize=11.5, textcolor="0.2")
    ax.text(0.25, 0.215, "learns BOTH", ha="center", va="center",
            fontsize=11, color="0.3", fontweight="bold")
    ax.text(0.25, 0.15, r"$p(x\,|\,c)$  and  $p(x\,|\,\varnothing)$",
            ha="center", va="center", fontsize=14, color="0.15")

    # ================= RIGHT PANEL: INFERENCE — steer with w =================
    ax.text(0.75, 0.95, "Inference — steer with w", ha="center", va="top",
            fontsize=14, color="0.25", fontweight="bold")

    # 1-D steering axis: null -> conditional -> guided
    x0, x1 = 0.58, 0.96
    yax = 0.62
    ax.annotate("", xy=(x1, yax), xytext=(x0, yax),
                arrowprops=dict(arrowstyle="-|>", color="0.55", lw=1.6))
    xnull, xcond = 0.62, 0.78
    xguided = 0.92
    # null anchor
    ax.plot([xnull], [yax], "o", ms=11, color=NULL, zorder=3)
    ax.text(xnull, yax + 0.045, "∅", ha="center", va="bottom",
            fontsize=15, color=NULL, fontweight="bold")
    ax.text(xnull, yax - 0.05, "unconditional", ha="center", va="top",
            fontsize=9.5, color=NULL)
    # conditional point
    ax.plot([xcond], [yax], "o", ms=11, color=COND, zorder=3)
    ax.text(xcond, yax + 0.045, "c", ha="center", va="bottom",
            fontsize=14, color=COND, fontweight="bold", style="italic")
    ax.text(xcond, yax - 0.05, "conditional (w=0)", ha="center", va="top",
            fontsize=9.5, color=COND)
    # guided point — pushed past c
    ax.plot([xguided], [yax], "o", ms=11, color=GOLD, zorder=3)
    ax.text(xguided, yax + 0.045, "guided", ha="center", va="bottom",
            fontsize=11, color=GOLD, fontweight="bold")

    # the "push past c by w" bracket from c to guided
    ax.annotate("", xy=(xguided, yax - 0.105), xytext=(xcond, yax - 0.105),
                arrowprops=dict(arrowstyle="-|>", color=GOLD, lw=2.2))
    ax.text((xcond + xguided) / 2, yax - 0.155, "push by  w", ha="center",
            va="top", fontsize=11, color=GOLD, fontweight="bold")

    # the equation
    ax.text(0.75, 0.31,
            r"$\ell_{\mathrm{guided}} \;=\; \ell(x\,|\,c)\;+\;w\,\left[\,\ell(x\,|\,c)-\ell(x\,|\,\varnothing)\,\right]$",
            ha="center", va="center", fontsize=15.5, color="0.1")

    # w readings
    ax.text(0.75, 0.185, "w = 0  →  plain conditional", ha="center", va="center",
            fontsize=11.5, color=COND)
    ax.text(0.75, 0.125, "w ↑  →  stronger class enforcement", ha="center",
            va="center", fontsize=11.5, color=GOLD, fontweight="bold")

    fig.savefig(FIG, dpi=200, bbox_inches="tight"); plt.close(fig)
    with Image.open(FIG) as im:
        print(f"fig: {FIG}  {im.size[0]}x{im.size[1]}px")


def fit(img, slot):
    sl, st, sw, sh = slot
    with Image.open(img) as im: iw, ih = im.size
    r, sr = iw/ih, sw/sh
    w, h = (sw, sw/r) if r > sr else (sh*r, sh)
    return sl+(sw-w)/2, st+(sh-h)/2, w, h


def set_title(tf, t):
    p0 = tf.paragraphs[0]
    if p0.runs:
        p0.runs[0].text = t
        for e in p0.runs[1:]: e._r.getparent().remove(e._r)
    else: p0.text = t
    for ep in list(tf.paragraphs[1:]): ep._p.getparent().remove(ep._p)


def set_annot(shape, text):
    tf = shape.text_frame; tf.word_wrap = True; tf.auto_size = MSO_AUTO_SIZE.NONE
    lines = text.split("\n"); tf.paragraphs[0].text = lines[0]
    for ep in list(tf.paragraphs[1:]): ep._p.getparent().remove(ep._p)
    for ln in lines[1:]: tf.add_paragraph().text = ln
    for p in tf.paragraphs:
        for r in p.runs: r.font.size = Pt(ANNOT_PT)
    shape.height = Inches(min(len(lines)*ANNOT_PT*1.2/72 + 0.15, 7.4 - Emu(shape.top).inches))


def main():
    if list(Path(PPTX).parent.glob("~$*.pptx")): raise SystemExit("ABORT: PowerPoint open.")
    build_fig()
    shutil.copy2(PPTX, f"/tmp/dplm_pptx_backup_{datetime.datetime.now():%Y%m%d_%H%M%S}.pptx")
    prs = Presentation(PPTX); tmpl = prs.slides[TEMPLATE_SLIDE_INDEX]
    t = prs.slides.add_slide(prs.slide_layouts[LAYOUT_INDEX])
    for ph in list(t.placeholders): ph.element.getparent().remove(ph.element)
    for sh in tmpl.shapes:
        if sh.shape_type == 13: continue
        t.shapes._spTree.insert_element_before(deepcopy(sh.element), "p:extLst")
    l, tp, w, h = fit(FIG, PIC_SLOT)
    t.shapes.add_picture(FIG, Inches(l), Inches(tp), Inches(w), Inches(h))
    for sh in t.shapes:
        if sh.shape_type == 13 or not sh.has_text_frame: continue
        if sh.name == TITLE_SHAPE: set_title(sh.text_frame, TITLE)
        elif sh.name == NUM_SHAPE: set_title(sh.text_frame, "")
        elif sh.name == ANNOT_SHAPE: set_annot(sh, ANNOT)
    prs.save(PPTX)
    print(f"appended CFG null-class slide at #{len(prs.slides)}; total {len(prs.slides)}")


if __name__ == "__main__":
    main()
