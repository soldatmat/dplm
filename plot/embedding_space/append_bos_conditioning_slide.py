#!/usr/bin/env python3
"""Append the BOS-conditioning-vectors slide to presentation/dplm.pptx (template
slide 232 layout). Run with PowerPoint CLOSED."""
from __future__ import annotations
import datetime, shutil
from copy import deepcopy
from pathlib import Path
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

ROOT = "/Users/soldatmat/Documents/terpene_synthases"
PPTX = f"{ROOT}/presentation/dplm.pptx"
FIG = f"{ROOT}/dplm/run/class_predictor/slide306_eval/bos_conditioning_map.png"
TEMPLATE_SLIDE_INDEX = 232; LAYOUT_INDEX = 1
TITLE_SHAPE = "TextovéPole 11"; NUM_SHAPE = "TextovéPole 12"; ANNOT_SHAPE = "TextBox 5"
PIC_SLOT = (0.12, 0.96, 8.48, 6.35)
TITLE = "Conditioning vectors in embedding space: MARTS-DB TPS BOS cloud + the 22 frozen per-class BOS-mean vectors"
ANNOT = ("Grey = all 1349 MARTS-DB\nTPS BOS embeddings\n(run_41V step-200000).\n\n"
         "Large coloured points =\nthe Embedding(22,640) the\nclass-conditioning\nClassEncoder is initialized\n"
         "to and FROZEN at (the\nper-class BOS mean),\ncoloured by first-\ncyclization class, labelled\nwith the class id.\n\n"
         "The 22 vectors are well-\nspread across the space\n(PCA std ~75% of the\ncloud's), i.e. the\nconditioning signal IS\ndistinct per class. So the\nweak steering (within-model\nΔ<=0.06) is the INJECTION\npathway (frozen, rank-1\nLoRA), not indistinguishable\nconditioning vectors.")


def fit(img, slot):
    sl, st, sw, sh = slot
    with Image.open(img) as im: iw, ih = im.size
    r, sr = iw/ih, sw/sh
    w, h = (sw, sw/r) if r > sr else (sh*r, sh)
    return sl+(sw-w)/2, st+(sh-h)/2, w, h


def set_text(tf, text):
    lines = text.split("\n"); p0 = tf.paragraphs[0]
    if p0.runs:
        p0.runs[0].text = lines[0]
        for e in p0.runs[1:]: e._r.getparent().remove(e._r)
    else: p0.text = lines[0]
    for ep in list(tf.paragraphs[1:]): ep._p.getparent().remove(ep._p)
    for ln in lines[1:]: tf.add_paragraph().text = ln


def main():
    if list(Path(PPTX).parent.glob("~$*.pptx")): raise SystemExit("ABORT: PowerPoint lock present — close it.")
    if not Path(FIG).exists(): raise SystemExit(f"missing {FIG}")
    bak = f"/tmp/dplm_pptx_backup_{datetime.datetime.now():%Y%m%d_%H%M%S}.pptx"
    shutil.copy2(PPTX, bak); print("backup ->", bak)
    prs = Presentation(PPTX); tmpl = prs.slides[TEMPLATE_SLIDE_INDEX]
    t = prs.slides.add_slide(prs.slide_layouts[LAYOUT_INDEX])
    for ph in list(t.placeholders): ph.element.getparent().remove(ph.element)
    for sh in tmpl.shapes:
        if sh.shape_type == 13: continue
        t.shapes._spTree.insert_element_before(deepcopy(sh.element), "p:extLst")
    l, tp, w, h = fit(FIG, PIC_SLOT); t.shapes.add_picture(FIG, Inches(l), Inches(tp), Inches(w), Inches(h))
    for sh in t.shapes:
        if sh.shape_type == 13 or not sh.has_text_frame: continue
        if sh.name == TITLE_SHAPE: set_text(sh.text_frame, TITLE)
        elif sh.name == NUM_SHAPE: set_text(sh.text_frame, "")
        elif sh.name == ANNOT_SHAPE: set_text(sh.text_frame, ANNOT)
    prs.save(PPTX); print(f"appended BOS-conditioning slide at #{len(prs.slides)}; total {len(prs.slides)}")


if __name__ == "__main__":
    main()
