#!/usr/bin/env python3
"""Redo slide #317 (the v2 de-overlapped BOS-conditioning map) so its cloud is the
slide-315 mean-pooled space with the BOS conditioning vectors projected in. Swaps
the picture + updates title/annotation on the existing slide (title-matched, robust
to index shifts). Run with PowerPoint CLOSED."""
from __future__ import annotations
import datetime, shutil
from pathlib import Path
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

ROOT = "/Users/soldatmat/Documents/terpene_synthases"
PPTX = f"{ROOT}/presentation/dplm.pptx"
FIG = f"{ROOT}/dplm/run/class_predictor/slide306_eval/bos_conditioning_map_meanspace.png"
TITLE_MATCH = "Conditioning vectors in slide-315 space (v2"
TITLE_SHAPE = "TextovéPole 11"; ANNOT_SHAPE = "TextBox 5"
PIC_SLOT = (0.12, 0.96, 8.48, 6.35)
NEW_TITLE = ("Conditioning vectors in slide-315 space (v2, labels de-overlapped): mean-pooled MARTS-DB TPS cloud "
             "+ the 22 frozen per-class BOS-mean vectors projected in")
NEW_ANNOT = ("Same cloud & PCA as\nslide 315 (mean-pooled\nMARTS-DB embeddings,\nrun_41V step-200000).\n\n"
             "The 22 conditioning\nvectors (BOS-mean, the\nEmbedding(22,640) the\nClassEncoder is frozen at)\n"
             "are standardized by the\nmean-cloud stats and\nprojected onto the mean-\npooled PCA axes; the\n"
             "t-SNE is recomputed on\nthe cloud + the 22 vectors\ntogether.\n\n"
             "NB the vectors live in\nBOS space, shown here on\nmean-pooled axes for an\napples-to-apples overlay\n"
             "with slide 315.\n\n"
             "Projected in, they cluster\nin a compact upper-central\nregion of the manifold\n(sterol class 12 the lone\noutlier).")


def fit(img, slot):
    sl, st, sw, sh = slot
    with Image.open(img) as im: iw, ih = im.size
    r, sr = iw/ih, sw/sh
    w, h = (sw, sw/r) if r > sr else (sh*r, sh)
    return sl+(sw-w)/2, st+(sh-h)/2, w, h


def set_text(tf, text):
    lines = text.split("\n"); p0 = tf.paragraphs[0]
    if p0.runs:
        p0.runs[0].text = lines[0]
        for e in p0.runs[1:]: e._r.getparent().remove(e._r)
    else: p0.text = lines[0]
    for ep in list(tf.paragraphs[1:]): ep._p.getparent().remove(ep._p)
    for ln in lines[1:]: tf.add_paragraph().text = ln


def main():
    if list(Path(PPTX).parent.glob("~$*.pptx")): raise SystemExit("ABORT: PowerPoint lock present — close it.")
    if not Path(FIG).exists(): raise SystemExit(f"missing {FIG}")
    bak = f"/tmp/dplm_pptx_backup_{datetime.datetime.now():%Y%m%d_%H%M%S}.pptx"
    shutil.copy2(PPTX, bak); print("backup ->", bak)
    prs = Presentation(PPTX)
    target = None
    for i, s in enumerate(prs.slides):
        for sh in s.shapes:
            if sh.has_text_frame and TITLE_MATCH in (sh.text_frame.text or ""):
                target = s; print(f"matched slide idx {i} (#{i+1})"); break
        if target: break
    if target is None: raise SystemExit(f"no slide titled like {TITLE_MATCH!r}")

    # remove existing picture(s)
    removed = 0
    for sh in list(target.shapes):
        if sh.shape_type == 13:
            sh.element.getparent().remove(sh.element); removed += 1
    print(f"removed {removed} old picture(s)")
    l, tp, w, h = fit(FIG, PIC_SLOT); target.shapes.add_picture(FIG, Inches(l), Inches(tp), Inches(w), Inches(h))

    for sh in target.shapes:
        if sh.shape_type == 13 or not sh.has_text_frame: continue
        if sh.name == TITLE_SHAPE: set_text(sh.text_frame, NEW_TITLE)
        elif sh.name == ANNOT_SHAPE: set_text(sh.text_frame, NEW_ANNOT)
    prs.save(PPTX); print("saved; slide #317 redone in mean-pooled (slide-315) space")


if __name__ == "__main__":
    main()
