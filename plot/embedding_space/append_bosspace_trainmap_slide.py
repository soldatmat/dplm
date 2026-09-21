#!/usr/bin/env python3
"""Append ONE slide: BOS-embedding-space version of slide #316 — all 1349
MARTS-DB training TPSs in the DPLM-150m BOS embedding space, PCA / t-SNE / UMAP,
coloured by first-cyclization class. Appends as a NEW slide (does NOT modify
slide #316/#317/#309). Run with PowerPoint CLOSED. Modeled on
append_baseline_overlay_slide.py."""
from __future__ import annotations
import datetime, shutil
from copy import deepcopy
from pathlib import Path
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

ROOT = "/Users/soldatmat/Documents/terpene_synthases"
PPTX = f"{ROOT}/presentation/dplm.pptx"
FIG = (f"{ROOT}/dplm/run/class_predictor/slide306_eval/"
       "training_tps_bosspace_pca_tsne_umap.png")
TEMPLATE_SLIDE_INDEX = 232; LAYOUT_INDEX = 1
TITLE_SHAPE = "TextovéPole 11"; NUM_SHAPE = "TextovéPole 12"; ANNOT_SHAPE = "TextBox 5"
PIC_SLOT = (0.12, 0.96, 8.48, 6.35)

TITLE = ("Training TPSs in the BOS embedding space – PCA / t-SNE / UMAP, "
         "coloured by first-cyclization class")
ANNOT = ("All 1349 MARTS-DB\ntraining TPSs in the\nDPLM-150m BOS\n"
         "embedding space\n(run_41V step-200000),\ncoloured by first-\n"
         "cyclization class\n(hue = substrate type).\n\n"
         "Same TPSs as slide #316\nbut BOS rather than\nsequence-mean pooling,\n"
         "plus a UMAP panel.\n\n"
         "z-score per dim, PCA\nvia SVD (PC1 23.5%,\nPC2 17.8% var); t-SNE\n"
         "perplexity 30, PCA-50\ninit; UMAP n_neighbors\n15, min_dist 0.1\n"
         "(all fit on the BOS cloud).")


def fit(img, slot):
    sl, st, sw, sh = slot
    with Image.open(img) as im: iw, ih = im.size
    r, sr = iw/ih, sw/sh
    w, h = (sw, sw/r) if r > sr else (sh*r, sh)
    return sl+(sw-w)/2, st+(sh-h)/2, w, h


def set_text(tf, text):
    lines = text.split("\n"); p0 = tf.paragraphs[0]
    if p0.runs:
        p0.runs[0].text = lines[0]
        for e in p0.runs[1:]: e._r.getparent().remove(e._r)
    else: p0.text = lines[0]
    for ep in list(tf.paragraphs[1:]): ep._p.getparent().remove(ep._p)
    for ln in lines[1:]: tf.add_paragraph().text = ln


def main():
    if list(Path(PPTX).parent.glob("~$*.pptx")):
        raise SystemExit("ABORT: PowerPoint lock present — close PowerPoint. "
                         f"PNG ready to append: {FIG}")
    if not Path(FIG).exists():
        raise SystemExit(f"missing {FIG}")
    bak = f"/tmp/dplm_pptx_backup_{datetime.datetime.now():%Y%m%d_%H%M%S}.pptx"
    shutil.copy2(PPTX, bak); print("backup ->", bak)

    prs = Presentation(PPTX)
    tmpl = prs.slides[TEMPLATE_SLIDE_INDEX]
    t = prs.slides.add_slide(prs.slide_layouts[LAYOUT_INDEX])
    for ph in list(t.placeholders):
        ph.element.getparent().remove(ph.element)
    for sh in tmpl.shapes:
        if sh.shape_type == 13: continue
        t.shapes._spTree.insert_element_before(deepcopy(sh.element), "p:extLst")
    l, tp, w, h = fit(FIG, PIC_SLOT)
    t.shapes.add_picture(FIG, Inches(l), Inches(tp), Inches(w), Inches(h))
    for sh in t.shapes:
        if sh.shape_type == 13 or not sh.has_text_frame: continue
        if sh.name == TITLE_SHAPE: set_text(sh.text_frame, TITLE)
        elif sh.name == NUM_SHAPE: set_text(sh.text_frame, "")
        elif sh.name == ANNOT_SHAPE: set_text(sh.text_frame, ANNOT)

    prs.save(PPTX)
    print(f"appended BOS-space trainmap slide at #{len(prs.slides)}; "
          f"total {len(prs.slides)}")


if __name__ == "__main__":
    main()
