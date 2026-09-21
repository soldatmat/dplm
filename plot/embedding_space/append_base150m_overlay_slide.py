#!/usr/bin/env python3
"""Append ONE slide showing the PRETRAINED DPLM-150m base model embedding-space
overlay (PCA | t-SNE | UMAP) to presentation/dplm.pptx (template slide 232).
Slide-342 twin for run_41V's un-finetuned starting point.
Run with PowerPoint CLOSED. Modeled on append_baseline_overlay_slide.py."""
from __future__ import annotations
import datetime, shutil
from copy import deepcopy
from pathlib import Path
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

ROOT = "/Users/soldatmat/Documents/terpene_synthases"
PPTX = f"{ROOT}/presentation/dplm.pptx"
FIG = (f"{ROOT}/dplm/plot/embedding_space/base150m_overlay_pca_tsne_umap.png")
TEMPLATE_SLIDE_INDEX = 232; LAYOUT_INDEX = 1
TITLE_SHAPE = "TextovéPole 11"; NUM_SHAPE = "TextovéPole 12"; ANNOT_SHAPE = "TextBox 5"
PIC_SLOT = (0.12, 0.96, 8.48, 6.35)

TITLE = ("Pretrained DPLM-150m base (run_41V's un-finetuned start) — 50 generated "
         "seqs in the MARTS-DB TPS embedding space (PCA / t-SNE / UMAP; space fit "
         "on train, generated projected in)")
ANNOT = ("Grey = all 1349 MARTS-DB\nTPSs (mean-pooled,\nrun_41V step-200000).\n\n"
         "Orange = 50 seqs from the\nRAW pretrained\nairkingbd/dplm_150m\n"
         "(no TPS fine-tuning),\nsame gen recipe as the\nrun_41V baseline\n"
         "(L=350, T=1.0,\ngumbel_argmax, 500 iter).\n\n"
         "Same train-only basis as\nslide 342 (PC1 41.4%,\nPC2 17.3% var).\n\n"
         "Surprise: base-gen does\nNOT land far out — in PCA\nit sits ON the train\n"
         "cloud (median radius 14.3\nvs train 15.8, 0.9x).\n\n"
         "But in t-SNE / UMAP it\nforms a few TIGHT, mostly\nSEPARATE sub-clusters\n"
         "(incl. an off-island at\nbottom) — protein-like &\noverlapping TPSs in coarse\n"
         "linear composition, yet\nlocally distinct from real\nTPSs. Fine-tuning\n"
         "diversifies/relocates onto\nthe manifold (cf. slide 342).")


def fit(img, slot):
    sl, st, sw, sh = slot
    with Image.open(img) as im: iw, ih = im.size
    r, sr = iw/ih, sw/sh
    w, h = (sw, sw/r) if r > sr else (sh*r, sh)
    return sl+(sw-w)/2, st+(sh-h)/2, w, h


def set_text(tf, text):
    lines = text.split("\n"); p0 = tf.paragraphs[0]
    if p0.runs:
        p0.runs[0].text = lines[0]
        for e in p0.runs[1:]: e._r.getparent().remove(e._r)
    else: p0.text = lines[0]
    for ep in list(tf.paragraphs[1:]): ep._p.getparent().remove(ep._p)
    for ln in lines[1:]: tf.add_paragraph().text = ln


def main():
    if list(Path(PPTX).parent.glob("~$*.pptx")):
        raise SystemExit("ABORT: PowerPoint lock present — close PowerPoint.")
    if not Path(FIG).exists():
        raise SystemExit(f"missing {FIG}")
    bak = f"/tmp/dplm_pptx_backup_{datetime.datetime.now():%Y%m%d_%H%M%S}.pptx"
    shutil.copy2(PPTX, bak); print("backup ->", bak)

    prs = Presentation(PPTX)
    tmpl = prs.slides[TEMPLATE_SLIDE_INDEX]
    t = prs.slides.add_slide(prs.slide_layouts[LAYOUT_INDEX])
    for ph in list(t.placeholders):
        ph.element.getparent().remove(ph.element)
    for sh in tmpl.shapes:
        if sh.shape_type == 13: continue
        t.shapes._spTree.insert_element_before(deepcopy(sh.element), "p:extLst")
    l, tp, w, h = fit(FIG, PIC_SLOT)
    t.shapes.add_picture(FIG, Inches(l), Inches(tp), Inches(w), Inches(h))
    for sh in t.shapes:
        if sh.shape_type == 13 or not sh.has_text_frame: continue
        if sh.name == TITLE_SHAPE: set_text(sh.text_frame, TITLE)
        elif sh.name == NUM_SHAPE: set_text(sh.text_frame, "")
        elif sh.name == ANNOT_SHAPE: set_text(sh.text_frame, ANNOT)

    prs.save(PPTX)
    print(f"appended base150m-overlay slide at #{len(prs.slides)}; "
          f"total {len(prs.slides)}")


if __name__ == "__main__":
    main()
