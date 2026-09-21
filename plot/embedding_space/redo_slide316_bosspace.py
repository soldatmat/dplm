#!/usr/bin/env python3
"""Redo slide #316 (the BOS-conditioning-vectors map) so BOTH the grey cloud AND
the 22 per-class conditioning vectors live in the SAME BOS embedding space —
an honest same-representation plot, with a UMAP panel added (PCA | t-SNE | UMAP).

Swaps the picture + rewrites title/annotation on the existing slide, matched by
the title substring "Conditioning vectors in" (robust to index shifts). Backs up
to /tmp first; aborts if PowerPoint is open. Run with PowerPoint CLOSED."""
from __future__ import annotations
import datetime, shutil
from pathlib import Path
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

ROOT = "/Users/soldatmat/Documents/terpene_synthases"
PPTX = f"{ROOT}/presentation/dplm.pptx"
FIG = f"{ROOT}/dplm/run/class_predictor/slide306_eval/bos_conditioning_map_bosspace_3panel.png"
TITLE_MATCH = "Conditioning vectors in"
TITLE_SHAPE = "TextovéPole 11"; ANNOT_SHAPE = "TextBox 5"; NUM_SHAPE = "TextovéPole 12"
PIC_SLOT = (0.12, 0.96, 8.48, 6.35)
NEW_TITLE = ("Per-class BOS-mean conditioning vectors IN the BOS embedding space — "
             "PCA / t-SNE / UMAP (space fit on train MARTS-DB)")
NEW_ANNOT = ("Honest same-space plot:\nboth the grey cloud and\nthe 22 vectors are BOS\n"
             "embeddings (run_41V\nstep-200000) of the 1349\nMARTS-DB training TPSs.\n\n"
             "The 22 conditioning\nvectors (the frozen\nClassEncoder init) ARE the\n"
             "per-class BOS-mean\ncentroids of the cloud\n(verified, max|diff|~7e-7).\n\n"
             "Space fit on the cloud:\nPCA train-only SVD +\nproject; UMAP train-fit +\n"
             "transform; t-SNE\nrecomputed on cloud +\nvectors together.\n\n"
             "So this directly shows\nwhether the class-mean\nprototypes are well-\n"
             "separated and whether\nthey sit in dense data\nregions vs low-density\n"
             "voids (the phantom-mean\nquestion). Being centroids,\nthey land inside\n"
             "their classes; PCA puts\nmost in a dense central\nlobe, with sterol (12)\n"
             "and mono (5) the outliers.")


def fit(img, slot):
    sl, st, sw, sh = slot
    with Image.open(img) as im:
        iw, ih = im.size
    r, sr = iw / ih, sw / sh
    w, h = (sw, sw / r) if r > sr else (sh * r, sh)
    return sl + (sw - w) / 2, st + (sh - h) / 2, w, h


def set_text(tf, text):
    lines = text.split("\n"); p0 = tf.paragraphs[0]
    if p0.runs:
        p0.runs[0].text = lines[0]
        for e in p0.runs[1:]:
            e._r.getparent().remove(e._r)
    else:
        p0.text = lines[0]
    for ep in list(tf.paragraphs[1:]):
        ep._p.getparent().remove(ep._p)
    for ln in lines[1:]:
        tf.add_paragraph().text = ln


def main():
    if list(Path(PPTX).parent.glob("~$*.pptx")):
        raise SystemExit("ABORT: PowerPoint lock present — close it.")
    if not Path(FIG).exists():
        raise SystemExit(f"missing {FIG}")
    bak = f"/tmp/dplm_pptx_backup_{datetime.datetime.now():%Y%m%d_%H%M%S}.pptx"
    shutil.copy2(PPTX, bak); print("backup ->", bak)

    prs = Presentation(PPTX)
    target = None
    for i, s in enumerate(prs.slides):
        for sh in s.shapes:
            if sh.has_text_frame and TITLE_MATCH in (sh.text_frame.text or ""):
                target = s; print(f"matched slide idx {i} (#{i+1})"); break
        if target:
            break
    if target is None:
        raise SystemExit(f"no slide titled like {TITLE_MATCH!r}")

    removed = 0
    for sh in list(target.shapes):
        if sh.shape_type == 13:
            sh.element.getparent().remove(sh.element); removed += 1
    print(f"removed {removed} old picture(s)")
    l, tp, w, h = fit(FIG, PIC_SLOT)
    target.shapes.add_picture(FIG, Inches(l), Inches(tp), Inches(w), Inches(h))

    for sh in target.shapes:
        if sh.shape_type == 13 or not sh.has_text_frame:
            continue
        if sh.name == TITLE_SHAPE:
            set_text(sh.text_frame, NEW_TITLE)
        elif sh.name == ANNOT_SHAPE:
            set_text(sh.text_frame, NEW_ANNOT)
        elif sh.name == NUM_SHAPE:
            set_text(sh.text_frame, "")
    prs.save(PPTX)
    print("saved; slide #316 redone — 22 conditioning vectors in BOS space, PCA/t-SNE/UMAP")


if __name__ == "__main__":
    main()
