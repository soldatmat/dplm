#!/usr/bin/env python3
"""Append a simple, minimalistic CFG (classifier-free guidance) explainer slide
to dplm.pptx. Figure = the guidance equation + term key + a w trade-off strip.
PowerPoint must be CLOSED. Run with /opt/miniconda3/bin/python3."""
import datetime, shutil
from copy import deepcopy
from pathlib import Path
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.enum.text import MSO_AUTO_SIZE

ROOT = "/Users/soldatmat/Documents/terpene_synthases"
PPTX = f"{ROOT}/presentation/dplm.pptx"
FIG = f"{ROOT}/dplm/plot/cfg_explainer.png"
TEMPLATE_SLIDE_INDEX = 232; LAYOUT_INDEX = 1
TITLE_SHAPE = "TextovéPole 11"; NUM_SHAPE = "TextovéPole 12"; ANNOT_SHAPE = "TextBox 5"
PIC_SLOT = (0.12, 0.96, 8.48, 6.35); ANNOT_PT = 10.5

COND = "#4C72B0"    # conditional (target class)
NULL = "#888888"    # unconditional (learned null)
GOLD = "#E6B800"    # the w knob — matches the CFG colour used on slides #354/#359

TITLE = "Classifier-free guidance — a steering-strength knob"
ANNOT = ("CFG = one network, two modes.\n\n"
         "Training: 15% of steps drop the\n"
         "class label to a learned null ∅,\n"
         "so the model learns BOTH the\n"
         "conditional and unconditional\n"
         "distributions at once.\n\n"
         "Generation: blend them with w.\n"
         "• w = 0  → plain conditional\n"
         "  (faithful, weak steering)\n"
         "• w ↑  → amplifies the class\n"
         "  signal (stronger steering;\n"
         "  too high → off-manifold)\n\n"
         "w is inference-only: one trained\n"
         "model, steering dialed at\n"
         "sampling time.")


def build_fig():
    fig = plt.figure(figsize=(8.5, 6.4))
    ax = fig.add_axes([0, 0, 1, 1]); ax.axis("off")
    ax.set_xlim(0, 1); ax.set_ylim(0, 1)

    # section label
    ax.text(0.5, 0.95, "Guided generation: combine two logit predictions",
            ha="center", va="top", fontsize=13, color="0.35")

    # main equation
    ax.text(0.5, 0.78,
            r"$\ell_{\mathrm{guided}} \;=\; (1+w)\,\ell(x\,|\,c)\;-\;w\,\ell(x\,|\,\varnothing)$"
            .replace(r"\varnothing", r"\emptyset"),
            ha="center", va="center", fontsize=27, color="0.1")

    # rearranged (the "steering" reading)
    ax.text(0.5, 0.605, "equivalently", ha="center", va="center",
            fontsize=11, color="0.5", style="italic")
    ax.text(0.5, 0.515,
            r"$=\; \ell(x\,|\,c)\;+\;w\,\left[\,\ell(x\,|\,c)-\ell(x\,|\,\emptyset)\,\right]$",
            ha="center", va="center", fontsize=20, color="0.1")
    ax.text(0.5, 0.43, "conditional  +  w × (class signal)",
            ha="center", va="center", fontsize=11, color=GOLD, fontweight="bold")

    # term key
    ax.text(0.18, 0.305, r"$\ell(x\,|\,c)$", ha="left", va="center",
            fontsize=14, color=COND, fontweight="bold")
    ax.text(0.34, 0.305, "conditional logits (target class $c$)",
            ha="left", va="center", fontsize=12, color="0.2")
    ax.text(0.18, 0.245, r"$\ell(x\,|\,\emptyset)$", ha="left", va="center",
            fontsize=14, color=NULL, fontweight="bold")
    ax.text(0.34, 0.245, "unconditional logits (learned null)",
            ha="left", va="center", fontsize=12, color="0.2")
    ax.text(0.18, 0.185, r"$w$", ha="left", va="center",
            fontsize=14, color=GOLD, fontweight="bold")
    ax.text(0.34, 0.185, "guidance weight — steering strength",
            ha="left", va="center", fontsize=12, color="0.2")

    # w trade-off strip
    ax.annotate("", xy=(0.92, 0.075), xytext=(0.10, 0.075),
                arrowprops=dict(arrowstyle="-|>", color=GOLD, lw=2.4))
    ax.text(0.10, 0.115, "w = 0", ha="left", va="bottom", fontsize=11,
            color="0.25", fontweight="bold")
    ax.text(0.10, 0.035, "faithful, weak steering", ha="left", va="top",
            fontsize=9.5, color="0.45")
    ax.text(0.92, 0.115, "large w", ha="right", va="bottom", fontsize=11,
            color="0.25", fontweight="bold")
    ax.text(0.92, 0.035, "strong steering, off-manifold risk", ha="right",
            va="top", fontsize=9.5, color="0.45")

    fig.savefig(FIG, dpi=200, bbox_inches="tight"); plt.close(fig)
    with Image.open(FIG) as im:
        print(f"fig: {FIG}  {im.size[0]}x{im.size[1]}px")


def fit(img, slot):
    sl, st, sw, sh = slot
    with Image.open(img) as im: iw, ih = im.size
    r, sr = iw/ih, sw/sh
    w, h = (sw, sw/r) if r > sr else (sh*r, sh)
    return sl+(sw-w)/2, st+(sh-h)/2, w, h


def set_title(tf, t):
    p0 = tf.paragraphs[0]
    if p0.runs:
        p0.runs[0].text = t
        for e in p0.runs[1:]: e._r.getparent().remove(e._r)
    else: p0.text = t
    for ep in list(tf.paragraphs[1:]): ep._p.getparent().remove(ep._p)


def set_annot(shape, text):
    tf = shape.text_frame; tf.word_wrap = True; tf.auto_size = MSO_AUTO_SIZE.NONE
    lines = text.split("\n"); tf.paragraphs[0].text = lines[0]
    for ep in list(tf.paragraphs[1:]): ep._p.getparent().remove(ep._p)
    for ln in lines[1:]: tf.add_paragraph().text = ln
    for p in tf.paragraphs:
        for r in p.runs: r.font.size = Pt(ANNOT_PT)
    shape.height = Inches(min(len(lines)*ANNOT_PT*1.2/72 + 0.15, 7.4 - Emu(shape.top).inches))


def main():
    if list(Path(PPTX).parent.glob("~$*.pptx")): raise SystemExit("ABORT: PowerPoint open.")
    build_fig()
    shutil.copy2(PPTX, f"/tmp/dplm_pptx_backup_{datetime.datetime.now():%Y%m%d_%H%M%S}.pptx")
    prs = Presentation(PPTX); tmpl = prs.slides[TEMPLATE_SLIDE_INDEX]
    t = prs.slides.add_slide(prs.slide_layouts[LAYOUT_INDEX])
    for ph in list(t.placeholders): ph.element.getparent().remove(ph.element)
    for sh in tmpl.shapes:
        if sh.shape_type == 13: continue
        t.shapes._spTree.insert_element_before(deepcopy(sh.element), "p:extLst")
    l, tp, w, h = fit(FIG, PIC_SLOT)
    t.shapes.add_picture(FIG, Inches(l), Inches(tp), Inches(w), Inches(h))
    for sh in t.shapes:
        if sh.shape_type == 13 or not sh.has_text_frame: continue
        if sh.name == TITLE_SHAPE: set_title(sh.text_frame, TITLE)
        elif sh.name == NUM_SHAPE: set_title(sh.text_frame, "")
        elif sh.name == ANNOT_SHAPE: set_annot(sh, ANNOT)
    prs.save(PPTX)
    print(f"appended CFG explainer slide at #{len(prs.slides)}; total {len(prs.slides)}")


if __name__ == "__main__":
    main()
